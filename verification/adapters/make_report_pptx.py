"""보고용 PPT 1장 생성 (REPORT_GHG_GEARs_verification.md 요약). 실행: uv run --with python-pptx python make_report_pptx.py"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[2] / "verification/evidence/C02/manual_S04/GHG_GEARs_verification_1p.pptx"
NAVY, GREY, RED, AMB, GRN = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x59, 0x59, 0x59), RGBColor(0xC0, 0x39, 0x2B), RGBColor(0xD6, 0x8A, 0x00), RGBColor(0x2E, 0x7D, 0x32)

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])


def text(x, y, w, h, lines, size=10, bold_first=False, color=GREY, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = ln; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold_first and i == 0
        p.space_after = Pt(2)
    return tb


def table(x, y, w, rows, col_w, size=8.5, head=True):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.22 * len(rows))).table
    for j, cw in enumerate(col_w): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = str(v); c.margin_left = c.margin_right = Inches(0.04); c.margin_top = c.margin_bottom = Inches(0.01)
            for p in c.text_frame.paragraphs:
                for r in p.runs: r.font.size = Pt(size); r.font.bold = head and i == 0; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if head and i == 0 else GREY
            if head and i == 0: c.fill.solid(); c.fill.fore_color.rgb = NAVY
    return t


# 제목
text(0.4, 0.2, 12.5, 0.5, ["K-MDS Use Case #3  선박 환경규제 의무보고 데이터 상호운용성 검증 결과 — IDS Consumer → IMO Compendium 표준 매핑 → KR GEARs Voyage Template/API"], 16, True, NAVY)
text(0.4, 0.62, 12.5, 0.3, ["RS-2024-00454634 3차년도 · case C02 · 2026-09-13 · 입력: vessellink(랩오투원 LAB021) Noon Report API, IDS 경유 수신본 = 원본(sha256 fb23ffbd…) · 선박 IMO 00000009(시뮬레이터) 2026-04-19~30, 이벤트 12건/필드 71종"], 9, False, GREY)
# 흐름
text(0.4, 0.95, 12.5, 0.35, ["vessellink ─Provider API─▶ K-MDS Provider Connector ─IDS(DAPS·계약)─▶ KR Consumer Connector ─▶ ① IMO Compendium(FAL50) 매핑 ─▶ ② GEARs Type1 Voyage Template ─▶ ③ Nexawave Post DCS/MRV Voyage Template API(전송 미실행)"], 9.5, False, NAVY)

# 좌: 매핑
text(0.4, 1.35, 4.2, 0.3, ["① 표준 매핑 (Provider 71 필드 → IMO Compendium)"], 11, True, NAVY)
table(0.4, 1.65, 4.2, [["구분", "건수", "비고"],
                       ["매핑 완료", "70 / 71 (98.6%)", "미해결 isEuPort"],
                       ["코드북 1:1", "34", "voyNo→IMO0191 등"],
                       ["연료 구조 해소", "24", "IMO0654 + IMO0670/0893/0673/0903"],
                       ["최특정·컨텍스트", "9 + 3", "draft→0621/0622, 항구→0108/0111"],
                       ["이벤트 코드", "6종", "EV01/02/16/10, BUNKERING·CARGO→EV28"],
                       ["연료 코드", "6종", "Lsfo→VLSFO2020→HFO열 (FAL50 설명)"]], [1.3, 1.1, 1.8])
text(0.4, 3.3, 4.2, 0.9, ["S03에서 에이전트 단독 매핑 2~3/60 → 코드북 1:N(42필드)·후보집합 제한을 규칙으로 해소.", "규칙은 gears_voyage_transform.py 에 고정(재현 가능)."], 8.5)

# 중: GEARs 변환
text(4.8, 1.35, 4.4, 0.3, ["② GEARs Type1 변환 — 레그 1 (PACTB → GTPRQ)"], 11, True, NAVY)
table(4.8, 1.65, 4.4, [["Type1 열", "값", "근거"],
                       ["Voyage / 출발→도착", "1 / PACTB→GTPRQ", "IMO0191/0111/0108"],
                       ["출발 / 도착 (UTC)", "04-19 20:36 / 04-25 01:00", "IMO0065 / IMO0063"],
                       ["항해시간 / 거리", "124.4 h / 871.6 nm", "ATA−ATD / ΣIMO0613"],
                       ["HFO ROB 출/도, 소비, 수급", "298.69/841.90, 54.96, 599.97", "VLSFO2020 계열"],
                       ["MGO ROB 출/도, 소비, 수급", "115.98/202.26, 0.60, 88.30", "MGO+ULSMGO"],
                       ["화물작업(도착)", "Yes", "CARGO_WORK 이벤트"],
                       ["출항 화물작업·정박유휴·화물량", "미도출", "Provider 항목 없음"]], [1.55, 1.7, 1.15])
text(4.8, 3.5, 4.4, 0.7, ["레그 2(GTPRQ 출항 04-29, 337 nm)는 도착 이벤트 없음 → 미완료 레그.", "③ API 리스트: 132 필드 중 69 사용, 필수 52 중 50 충족(레그1). ImoNo 7자리 위반(테스트 ID). Token·endpoint 미확보로 전송 미실행."], 8.5)

# 우: 판정
text(9.4, 1.35, 3.6, 0.3, ["판정: PARTIAL"], 13, True, AMB)
table(9.4, 1.7, 3.6, [["검사 53건", "결과"],
                      ["PASS", "36"],
                      ["FAIL (필수 미도출)", "6 — 출항 화물작업, 정박유휴시간, MRV 화물량×2행"],
                      ["WARNING", "5 — MGO 질량수지 +1.42 MT, 거리 교차 871.6 vs 969.6, 정박 ROB 감소 19.07 MT 미보고, PACTB UNLOCODE 미수록"],
                      ["미완료 레그", "6 (레그 2 도착 전)"],
                      ["HFO 질량수지", "+1.80 MT (허용 ±4.21) 통과"],
                      ["재현성", "입력 sha256·registry 1205·스크립트 고정"]], [1.15, 2.45], 8)
text(9.4, 4.05, 3.6, 0.6, ["전달·표준매핑·GEARs 구조 변환은 성립. 단 Provider 이벤트만으로 DCS/MRV 필수 4종을 채울 수 없고 정박 소비량이 미보고 → 제출 데이터로는 미완성."], 8.5, color=GREY)

# 하단: 조치
text(0.4, 4.75, 12.5, 0.3, ["조치 요청 · 결정 사항"], 11, True, NAVY)
table(0.4, 5.05, 12.5, [["대상", "내용"],
                        ["랩오투원 (vessellink)", "FAL50 Event/Operation 코드 채택 · 출항 화물작업·앵커 이벤트·화물량 항목 추가 · 정박 이벤트에 연료 소비량 보고 · 결측 −9999/\"\" 정리 · 코드북 1:N(42/68) 컨텍스트 규칙 명시"],
                        ["KR GEARs / Nexawave", "Voyage Template API endpoint URL·Token 발급 · 템플릿 UNLOCODE 시트 갱신(PACTB 등) · 테스트 선박 IMO No 정책"],
                        ["KR 결정 (D6~D8)", "D6 VLSFO/ULSFO→HFO 열 분류 승인(FAL50 코드 설명 근거) · D7 VerificationTypeCode 002(IMO DCS) 단독/001(EU MRV) 병행 · D8 어댑터 규칙을 GHG AI Agent 후보집합·ingress에 반영(D2·D4)"]], [1.9, 10.6], 8.5)
text(0.4, 6.95, 12.5, 0.3, ["증적: k-mds/verification/evidence/C02/manual_S04/ (gears-transform/*, kr-systems-api/voyage_template_fields.json, reference/*, REPORT_GHG_GEARs_verification.md)"], 7.5, False, GREY)

prs.save(OUT); print("saved", OUT)
