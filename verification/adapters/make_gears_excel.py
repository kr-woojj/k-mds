"""매핑 결과물 엑셀 1부 — GEARs Type 1 Template/Sample 레이아웃으로 변환 행을 기록하고 근거 시트를 동봉.
입력: gears_voyage_transform.py 산출물(manual_S04/gears-transform/*), Noon/DAQ payload, 코드북, 템플릿.
출력: manual_S04/GEARs_Type1_Mapping_Result_<IMO>.xlsx
실행: uv run python make_gears_excel.py   (kr-ghg-ai-agent uv 환경: openpyxl)
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]  # k-mds 루트
S4 = ROOT / "verification/evidence/C02/manual_S04"
S3 = ROOT / "verification/evidence/C02/manual_S03"
TEMPLATE = ROOT / "data/raw/KR-Systems/GEARs_Template_Type1_Rev.2.1.xlsm"
GT = S4 / "gears-transform"

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F"); SUB_FILL = PatternFill("solid", fgColor="D9E1F2"); WARN_FILL = PatternFill("solid", fgColor="FFF2CC"); FAIL_FILL = PatternFill("solid", fgColor="F8CBAD")
THIN = Border(*(Side(style="thin", color="999999"),) * 4)
WHITE_B = Font(bold=True, color="FFFFFF"); BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def norm(s): return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def template_columns():
    ws = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)["Type 1"]
    rows = {r: [c for c in ws.iter_rows(min_row=r, max_row=r, values_only=True)][0] for r in (4, 5, 6, 7)}
    cols, grp, last = [], None, None
    for c in range(3, len(rows[5])):  # 템플릿 D열(index 3)부터 → 출력 A열 (Sample 과 동일 정렬)
        g, n, sb, f = (re.sub(r"\s+", " ", str(rows[r][c])).strip() if rows[r][c] is not None else None for r in (4, 5, 6, 7))
        if g: grp = g
        if n: last = n
        cols.append({"tcol": get_column_letter(c + 1), "group": grp, "name": n, "name_ff": last, "sub": sb, "format": f, "raw_name": rows[5][c], "raw_sub": rows[6][c]})
    while cols and not (cols[-1]["name"] or cols[-1]["sub"]): cols.pop()
    return cols


ROW_KEYS: dict[str, str] = {}  # norm(key) → key (gears-type1-rows.json 의 실제 키)


def _rk(candidate):
    """템플릿 표기(예: 'LPG(Butane)', 'LNG ROB at departure(MT)')와 행 키('LPG (Butane)', '… departure (MT)') 의 공백 차이를 정규화로 흡수."""
    return ROW_KEYS.get(norm(candidate))


def row_key_for(col):
    """Type1 열 → gears-type1-rows.json 키."""
    g, n, sb = col["group"] or "", col["name_ff"] or "", col["sub"]
    if "Bunkering before arrival" in g:
        if n == "Voyage Number": return "Voyage No."
        return _rk(f"Bunkering before arrival: Received Quantity ({sb})") if sb else None
    if n == "Cargo carried" and sb: return _rk(f"Cargo carried ({sb})")
    if n == "Transport Work" and sb: return _rk(f"Transport Work ({sb})")
    if n == "Cargo operation (YES/NO)": return None  # 위치로 결정
    return _rk(n) or (n if n in ROW_KEYS.values() else None)


def build_type1(wb, cols, rows, findings):
    ws = wb.active; ws.title = "Type 1"
    # 1행 그룹(병합), 2~3행 항목/연료, 4행~ 데이터 — GEARs_Sample_Type 1.xlsx 레이아웃
    c0 = 1
    for i, col in enumerate(cols):
        c = c0 + i
        ws.cell(2, c, col["raw_name"]); ws.cell(3, c, col["raw_sub"])
        for r in (2, 3):
            ws.cell(r, c).alignment = CENTER; ws.cell(r, c).border = THIN; ws.cell(r, c).font = BOLD; ws.cell(r, c).fill = SUB_FILL
        if col["raw_name"] and not col["raw_sub"] and not (i + 1 < len(cols) and cols[i + 1]["raw_sub"] and not cols[i + 1]["raw_name"]):
            ws.merge_cells(start_row=2, start_column=c, end_row=3, end_column=c)
        ws.column_dimensions[get_column_letter(c)].width = 13
    # 이름 가로 병합(연료 sub 가 이어지는 구간)
    i = 0
    while i < len(cols):
        j = i
        while j + 1 < len(cols) and cols[j + 1]["raw_name"] is None and cols[j + 1]["raw_sub"]: j += 1
        if j > i and cols[i]["raw_sub"]: ws.merge_cells(start_row=2, start_column=c0 + i, end_row=2, end_column=c0 + j)
        i = j + 1
    # 그룹 병합
    i = 0
    while i < len(cols):
        j = i
        while j + 1 < len(cols) and cols[j + 1]["group"] == cols[i]["group"]: j += 1
        ws.cell(1, c0 + i, cols[i]["group"]); ws.cell(1, c0 + i).fill = HEAD_FILL; ws.cell(1, c0 + i).font = WHITE_B; ws.cell(1, c0 + i).alignment = CENTER
        if j > i: ws.merge_cells(start_row=1, start_column=c0 + i, end_row=1, end_column=c0 + j)
        i = j + 1
    ws.row_dimensions[2].height = 48; ws.row_dimensions[1].height = 22
    # 데이터
    cargo_seen = 0
    fail_cells = {(f["row"], f["col"]): f for f in findings}
    for ri, row in enumerate(rows):
        r = 4 + ri; cargo_seen = 0
        for i, col in enumerate(cols):
            key = row_key_for(col)
            if col["name_ff"] == "Cargo operation (YES/NO)":
                key = "Cargo operation (DEPARTURE)" if cargo_seen == 0 else "Cargo operation (ARRIVAL)"; cargo_seen += 1
            if key is None: continue
            v = row.get(key)
            cell = ws.cell(r, c0 + i)
            if v is not None and re.search(r"Date \(UTC\)", key):
                cell.value = datetime.strptime(v, "%Y-%m-%d"); cell.number_format = "yyyy-mm-dd"
            elif v is not None and re.search(r"time \(UTC\)", key):
                cell.value = v
            else:
                cell.value = v
            cell.border = THIN
            f = fail_cells.get((f"row{ri+1}", key)) or fail_cells.get((f"row{ri+1}", key.split(" ROB")[0].split(" Consumption")[0] if "(MT)" in key else "\0"))
            if v is None and f: cell.fill = FAIL_FILL if f["severity"] == "FAIL" else WARN_FILL
        ws.cell(r, c0 + len(cols) + 1, "레그 미완료(도착 이벤트 없음)" if row.get("Port code (ARRIVAL)") is None else "완료 레그")
    ws.freeze_panes = "B4"
    return ws


def sheet_table(wb, title, header, rows, widths=None, fills=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(1, c).fill = HEAD_FILL; ws.cell(1, c).font = WHITE_B; ws.cell(1, c).alignment = CENTER; ws.cell(1, c).border = THIN
    for i, r in enumerate(rows):
        ws.append(list(r))
        if fills and fills[i]:
            for c in range(1, len(header) + 1): ws.cell(i + 2, c).fill = fills[i]
    for i, w in enumerate(widths or [], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def main():
    cols = template_columns()
    rows = json.loads((GT / "gears-type1-rows.json").read_text("utf-8"))
    ROW_KEYS.update({norm(k): k for k in rows[0].keys()})
    rep = json.loads((GT / "validation-report.json").read_text("utf-8"))
    mapping = json.loads((GT / "mapping-table.json").read_text("utf-8"))
    canon = json.loads((GT / "imo-canonical-events.json").read_text("utf-8"))
    api = json.loads((GT / "kr-systems-api-datalist.json").read_text("utf-8"))
    noon = json.loads((S3 / "step-04/payloads/Noon_Report_API__e5e3be7e7a31.json").read_text("utf-8"))["data"]
    daq = json.loads((S3 / "step-04/payloads/DAQ_Logger_API__7679054100f6.json").read_text("utf-8"))["data"]
    lcb = json.loads((S3 / "step-02/api/logger-code-book.txt").read_text("utf-8"))["data"]["items"]
    lby = {it["externalKey"]: it for it in lcb}
    general = noon["general"]; imo = general.get("imoNo")

    wb = openpyxl.Workbook()
    build_type1(wb, cols, rows, rep["findings"])

    # Type1 열 ← 변환값 ← IMO 요소 ← Provider 필드 (열 단위 매핑 근거)
    derive = {"Voyage No.": ("IMO0191", "voyNo"), "Port code (DEPARTURE)": ("IMO0111", "portCode@DEPARTURE_SBY"), "Departure Date (UTC)": ("IMO0065", "dateEventUtc@DEPARTURE_SBY"),
              "Departure time (UTC)": ("IMO0065", "dateEventUtc@DEPARTURE_SBY"), "Cargo operation (DEPARTURE)": ("—", "미도출(이전 정박 CARGO_WORK 이벤트 유무)"),
              "Port code (ARRIVAL)": ("IMO0108", "portCode@ARRIVAL"), "Arrival Date (UTC)": ("IMO0063", "dateEventUtc@ARRIVAL"), "Arrival time (UTC)": ("IMO0063", "dateEventUtc@ARRIVAL"),
              "Cargo operation (ARRIVAL)": ("IMO0597=EV28", "CARGO_WORK 이벤트 존재"), "Time spent at sea (hours)": ("IMO0063−IMO0065", "ATA−ATD"),
              "Distance travelled (nm)": ("Σ IMO0613", "steamingDistanceSlr 합(출항 후~도착)"), "Total idle spent time at anchorage (hours)": ("—", "미도출(앵커 이벤트 없음)"),
              "Cargo carried (Passenger)": ("—", "Provider 항목 없음"), "Cargo carried (Weight or Volume)": ("—", "Provider 항목 없음")}
    fuel_src = {"HFO": "robHfo+robLsfo+robUlsfo / consumption*{Hfo,Lsfo,Ulsfo}", "MDO": "robDo / consumption*Do", "MGO": "robMgo+robUlsmgo / consumption*{Mgo,Ulsmgo}"}
    fuel_imo = {"HFO": "IMO0654∈{HFO,VLSFO2020,ULSFO2020} + IMO0674 / IMO0670·0893·0673·0903", "MDO": "IMO0654=MDO + IMO0674 / …", "MGO": "IMO0654∈{MGO,ULSMGO2020} + IMO0674 / …"}
    trows = []
    for col in cols:
        key = row_key_for(col) or col["name_ff"]
        if key == "Cargo operation (YES/NO)": continue
        vals = [r.get(key) for r in rows] if key else [None, None]
        m = re.match(r"^(.+?) (ROB at departure|ROB at arrival)|^At Berth (.+?) Consumption|^(.+?) Consumption \(MT\)|Received Quantity \((.+?)\)", key or "")
        fuel = next((g for g in (m.groups() if m else []) if g and g in fuel_src), None) if m else None
        imo_ref, src = derive.get(key, (fuel_imo.get(fuel, "—") if fuel else "—", fuel_src.get(fuel, "해당 연료 없음(0)") if fuel else "—"))
        trows.append([col["tcol"], col["group"], key, col["format"], imo_ref, src, vals[0], vals[1] if len(vals) > 1 else None])
    sheet_table(wb, "Mapping_Type1←IMO←Noon", ["템플릿 열", "그룹", "Type1 항목", "형식", "IMO Compendium 근거", "Provider(vessellink) 원천 필드/산식", "레그1 값", "레그2 값"], trows, [8, 30, 42, 12, 34, 44, 14, 14])

    # Provider 필드 → IMO ID 표
    mrows = [[k, v["imo_id"] or "—", v["imo_name"] or "", v["format"] or "", v["rule"], v["note"] or "", v["seen"]] for k, v in sorted(mapping.items())]
    fills = [WARN_FILL if r[1] == "—" else None for r in mrows]
    sheet_table(wb, "Mapping_Noon→IMO", ["Provider 필드(externalKey)", "IMO Data Number", "FAL50 명칭", "FAL50 형식", "매핑 규칙", "비고", "출현 이벤트 수"], mrows, [26, 16, 40, 12, 24, 44, 12], fills)

    # IMO 정준 이벤트 (평면화)
    ids = sorted({k for e in canon for k in e if k.startswith("IMO")})
    crows = []
    for e in canon:
        cons = "; ".join(f"{fc}: " + ", ".join(f"{k}={v}" for k, v in occ.items() if k != "IMO0654") for fc, occ in e["consumption_by_fuel_type"].items())
        robs = "; ".join(f"{fc}={occ.get('IMO0674')}" for fc, occ in e["rob_by_fuel_type"].items())
        crows.append([e["_time"], e["_event_key"], e["IMO0597"]] + [e.get(i) for i in ids if i != "IMO0597"] + [cons, robs])
    sheet_table(wb, "IMO_Canonical_Events", ["시각(UTC)", "Provider eventKey", "IMO0597"] + [i for i in ids if i != "IMO0597"] + ["Consumption By Fuel Type(IMO0654+장비요소)", "ROB by fuel(IMO0674)"], crows, [22, 16, 10] + [12] * (len(ids) - 1) + [60, 40])

    # DAQ (ISO 19848) 매핑 + 범위 검사
    dkeys = sorted({k for it in daq["items"] for k in it})
    drows, dfills = [], []
    for k in dkeys:
        it = lby.get(k, {}); vals = [it2.get(k) for it2 in daq["items"] if isinstance(it2.get(k), (int, float))]
        lo, hi = it.get("minValue"), it.get("maxValue")
        oor = sum(1 for v in vals if v == -9999 or (lo is not None and hi is not None and hi > lo and not (lo <= v <= hi)))
        drows.append([k, it.get("id"), it.get("localId"), it.get("unit"), it.get("type"), lo, hi, it.get("description"), len(vals), oor, (min(vals) if vals else None), (max(vals) if vals else None)])
        dfills.append(WARN_FILL if oor else None)
    sheet_table(wb, "DAQ_ISO19848_Mapping", ["DAQ 필드(externalKey)", "KSDX ID", "ISO 19848 LocalId(Naming rule)", "단위", "타입", "min", "max", "설명", "값 개수", "범위 밖/−9999", "최소", "최대"], drows, [22, 10, 52, 8, 8, 8, 8, 34, 8, 12, 10, 10], dfills)

    # 검증 결과
    vrows = [[f["row"], f.get("col") or f.get("field"), f["code"], f["severity"], f.get("detail", "")] for f in rep["findings"]] + \
            [[f["row"], f.get("field"), f["code"], f["severity"], f.get("detail", "")] for f in api["validation"]["findings"]]
    vfills = [FAIL_FILL if r[3] == "FAIL" else WARN_FILL for r in vrows]
    ws = sheet_table(wb, "Validation", ["행", "열/필드", "코드", "심각도", "상세"], vrows, [8, 34, 32, 14, 80], vfills)
    ws.append([]); ws.append(["요약", f"템플릿 검사 {rep['checks']}건: FAIL {rep['fail']} / WARNING {rep['warning']} / 미완료 레그 {rep['incomplete']} / PASS {rep['checks'] - len(rep['findings'])}",
                             f"API 필드 {api['validation']['spec_fields']} (필수 {api['validation']['mandatory_fields']}), API 검증 FAIL {api['validation']['fail']}", "판정 PARTIAL"])

    # API 데이터 리스트
    props = list(api["rows"][0].keys())
    sheet_table(wb, "API_DataList", ["API 속성"] + [f"레그{i+1}" for i in range(len(api["rows"]))], [[p] + [r.get(p) for r in api["rows"]] for p in props], [26, 18, 18])

    # README
    r = wb.create_sheet("README", 0)
    lines = [["GEARs Type 1 매핑 결과물 (K-MDS Use Case #3, case C02, 2026-09-13)"], [],
             ["선박", f"IMO {imo} / {general.get('shipName')} (vessellink 시뮬레이터 선박, callsign {general.get('callsign')})"],
             ["입력", "vessellink(랩오투원 LAB021) Noon Report API 12 이벤트 / DAQ Logger API 20 레코드 — K-MDS IDS Consumer(KR) 경유 수신, 원본과 sha256 동일"],
             ["경로", "Provider payload → 코드북+FAL50 규칙으로 IMO Compendium 매핑 → GEARs Type 1 Voyage Template(Rev.2.1) 행 → Nexawave Post DCS/MRV Voyage Template API 속성"],
             ["시트", "Type 1: Sample 레이아웃(1행 그룹, 2~3행 항목/연료, 4행~ 데이터). 음영: 빨강=필수 미도출(FAIL), 노랑=경고"],
             ["", "Mapping_Type1←IMO←Noon: 템플릿 열별 IMO 근거·원천 필드 / Mapping_Noon→IMO: Provider 71필드 매핑 / IMO_Canonical_Events: 이벤트별 IMO 요소값"],
             ["", "DAQ_ISO19848_Mapping: DAQ 72필드 → KSDX ID·ISO 19848 LocalId, 범위 검사 / Validation: 검사 결과 / API_DataList: Nexawave 속성명 데이터"],
             ["연료 열 분류", "Lsfo→VLSFO2020, Ulsfo→ULSFO2020 → HFO 열 (FAL50 Fuel type 설명 'HFO w/ S 0.1~0.5%'); Do→MDO; Mgo·Ulsmgo→MGO — KR 승인 필요(D6)"],
             ["판정", f"PARTIAL — 템플릿 검사 {rep['checks']}건 FAIL {rep['fail']}·WARNING {rep['warning']}·미완료 {rep['incomplete']}; API 필수 52 중 레그1 50 충족"],
             ["재현", "k-mds/verification/adapters/gears_voyage_transform.py → make_gears_excel.py (registry FAL50 elementCount 1205, 코드북 sha 기록은 evidence 참조)"],
             ["미전송", "Nexawave API 는 Token·endpoint 미확보로 전송하지 않음. 값은 검토용."]]
    for ln in lines: r.append(ln)
    r.column_dimensions["A"].width = 14; r.column_dimensions["B"].width = 140; r["A1"].font = Font(bold=True, size=13)
    out = S4 / f"GEARs_Type1_Mapping_Result_IMO{imo}.xlsx"
    wb.save(out); print("saved", out)
    return out


if __name__ == "__main__":
    main()
