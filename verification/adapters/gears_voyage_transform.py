"""LAB021(vessellink) Noon Report payload → IMO Compendium(FAL50) 표준 매핑 → KR GEARs Type1 Voyage Template 변환·검증.

근거(모두 파일에서 읽음, 기억 인용 없음):
- 코드북: evidence/C02/manual_S03/step-02/api/noon-code-book.txt (Provider 공개, externalKey→IMO ID, eventName 컨텍스트)
- FAL50: kr-ghg-ai-agent/var/registry.sqlite3 (data_element: name/format/code_list) + data/raw/FAL50/IMO Compendium.xlsx (Code list, Structure)
- GEARs: data/raw/KR-Systems/GEARs_Template_Type1_Rev.2.1.xlsm (Type 1 헤더 4~7행, UNLOCODE 시트, Setting 연료표)
출력: <out>/mapping-table.json|.md, imo-canonical-events.json, gears-type1-rows.json|.csv, kr-systems-api-datalist.json, validation-report.json

# ponytail: 항차 레그 = DEPARTURE_SBY → 다음 ARRIVAL. 정박(anchorage) 유휴시간은 이벤트에 앵커 코드가 없어 미도출(None).
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\kr-dev")
CODEBOOK = ROOT / "k-mds/verification/evidence/C02/manual_S03/step-02/api/noon-code-book.txt"
PAYLOAD = ROOT / "k-mds/verification/evidence/C02/manual_S03/step-04/payloads/Noon_Report_API__e5e3be7e7a31.json"
REGISTRY = ROOT / "kr-ghg-ai-agent/var/registry.sqlite3"
FAL50_XLSX = ROOT / "k-mds/data/raw/FAL50/IMO Compendium.xlsx"
TEMPLATE = ROOT / "k-mds/data/raw/KR-Systems/GEARs_Template_Type1_Rev.2.1.xlsm"

# ── 1. 표준 매핑 규칙 (코드북 1:N 해소) ───────────────────────────────────────────────
# eventKey(Provider) → 코드북 eventName 컨텍스트
EVENT_CTX = {"ARRIVAL": "Arrival", "DEPARTURE_SBY": "Departure S/By", "NOON_AT_SEA": "Noon at sea", "RUP": "R/UP", "BUNKERING": "Bunkering"}
# 컨텍스트로도 못 가르는 키: 가장 특정(specific)한 요소 또는 FAL50 'Consumption By Fuel Type' 구조(IMO0654 + 장비별 요소) 선택
PICK = {"draftFore": "IMO0621", "draftAft": "IMO0622", "draftMid": "IMO0357", "shipCourse": "IMO0332", "trueWindDirection": "IMO0627",
        "eventKey": "IMO0597", "masterName": "IMO0580", "berthName": "IMO0548", "distanceToGo": "IMO0615", "isEuPort": None,
        "portCode": {"Arrival": "IMO0108", "Departure S/By": "IMO0111", "Bunkering": "IMO0666"},
        "portName": {"Arrival": "IMO0109", "Departure S/By": "IMO0112", "Bunkering": "IMO0667"},
        "dateEventUtc": {"Arrival": "IMO0063", "Departure S/By": "IMO0065"},  # 그 외 IMO0603
        "positionLat": "IMO0601", "positionLon": "IMO0602"}
ENGINE_ELEMENT = {"Me": "IMO0670", "Ge": "IMO0893", "Blr": "IMO0673", "Other": "IMO0903"}  # Consumption By Fuel Type 구조(Structure 시트)
# Provider 연료 접미사 → FAL50 'Fuel type' 코드 → GEARs Type1 연료 열. 근거: FAL50 Code list 설명
# (VLSFO2020="HFO w/ S 0.1~0.5%", ULSFO2020="HFO w/ S ≤0.1%") → 잔사유 계열은 GEARs HFO 열.
FUEL = {"Hfo": ("HFO", "HFO"), "Lsfo": ("VLSFO2020", "HFO"), "Ulsfo": ("ULSFO2020", "HFO"), "Do": ("MDO", "MDO"),
        "Mgo": ("MGO", "MGO"), "Ulsmgo": ("ULSMGO2020", "MGO")}
EVENT_CODE = {"ARRIVAL": "EV01", "DEPARTURE_SBY": "EV02", "NOON_AT_SEA": "EV16", "RUP": "EV10", "BUNKERING": "EV28", "CARGO_WORK": "EV28"}
EVENT_CODE_NOTE = {"BUNKERING": "FAL50 Event type 에 Bunkering 없음 → EV28 Other event (F-12)", "CARGO_WORK": "하역은 Operation type(OP11/OP15) 성격 → EV28 + OP 미지정 (F-12)",
                   "RUP": "R/UP(Rung Up) → EV10 Begin of sea passage 로 해석(추정)", "DEPARTURE_SBY": "S/By(Stand-by) 시각을 ATD 로 사용(추정)"}
GEARS_FUELS = ["HFO", "LFO", "MDO", "MGO", "LPG (Propane)", "LPG (Butane)", "Methanol", "Ethanol", "LNG", "Other"]
CONS_RE = re.compile(r"^consumption(Me|Ge|Blr|Other)(Hfo|Lsfo|Ulsfo|Do|Mgo|Ulsmgo)$")
ROB_RE = re.compile(r"^rob(Hfo|Lsfo|Ulsfo|Do|Mgo|Ulsmgo)$")


def load_codebook():
    items = json.loads(CODEBOOK.read_text("utf-8"))["data"]["items"]
    by = {}
    for it in items:
        for k in it.get("externalKey") or []:
            by.setdefault(k, []).append(it)
    return by


def load_registry():
    con = sqlite3.connect(REGISTRY)
    cols = [r[1] for r in con.execute("pragma table_info(data_element)")]
    return {d["imo_data_number"]: d for d in (dict(zip(cols, r)) for r in con.execute("select * from data_element"))}


def load_code_lists():
    ws = openpyxl.load_workbook(FAL50_XLSX, read_only=True)["Code list"]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            out.setdefault(str(r[0]).strip(), {})[str(r[1])] = r[2]
    return out


def map_field(key, value, event_key, by, reg):
    """한 필드 → (imo_id, rule, note). 코드북에 없으면 imo_id None."""
    cands = by.get(key, [])
    ctx = EVENT_CTX.get(event_key)
    if not cands:
        return None, "NOT_IN_CODEBOOK", None
    ids = [c["id"] for c in cands]
    if len(ids) == 1:
        return ids[0], "CODEBOOK_1_1", None
    m = CONS_RE.match(key)
    if m:
        return ENGINE_ELEMENT[m.group(1)], "FUEL_TYPE_STRUCTURE", f"IMO0654={FUEL[m.group(2)][0]} occurrence"
    pick = PICK.get(key)
    if isinstance(pick, dict):
        chosen = pick.get(ctx)
        if chosen:
            return chosen, "EVENT_CONTEXT", f"eventName={ctx}"
        if key == "dateEventUtc":
            return "IMO0603", "EVENT_CONTEXT_DEFAULT", "Ship reporting date time"
        return None, "AMBIGUOUS_NO_CONTEXT", f"candidates={ids} ctx={ctx}"
    if pick:
        return pick, "MOST_SPECIFIC", f"candidates={ids}"
    narrowed = [c["id"] for c in cands if ctx and ctx in (c.get("eventName") or [])]
    if len(narrowed) == 1:
        return narrowed[0], "EVENT_CONTEXT", f"eventName={ctx}"
    return None, "AMBIGUOUS", f"candidates={ids}"


def canonicalize(payload, by, reg, code_lists):
    general = payload["general"]; events = sorted(payload["events"], key=lambda e: e["dateEventUtc"])
    table = {}; canon = []
    fuel_codes = code_lists["Fuel type"]
    for i, e in enumerate(events):
        rec = {"_event_key": e["eventKey"], "_time": e["dateEventUtc"], "IMO0140": general.get("imoNo"), "IMO0142": general.get("shipName"),
               "IMO0597": EVENT_CODE.get(e["eventKey"]), "consumption_by_fuel_type": {}, "rob_by_fuel_type": {}}
        for key, val in e.items():
            imo, rule, note = map_field(key, val, e["eventKey"], by, reg)
            row = table.setdefault(key, {"imo_id": imo, "rule": rule, "note": note, "imo_name": reg.get(imo, {}).get("name") if imo else None,
                                         "format": reg.get(imo, {}).get("format_spec") if imo else None, "seen": 0})
            row["seen"] += 1
            if imo is None:
                continue
            m = CONS_RE.match(key)
            if m:
                fc = FUEL[m.group(2)][0]
                rec["consumption_by_fuel_type"].setdefault(fc, {"IMO0654": fc})[imo] = val
                continue
            m = ROB_RE.match(key)
            if m:
                rec["rob_by_fuel_type"].setdefault(FUEL[m.group(1)][0], {"IMO0654": FUEL[m.group(1)][0]})["IMO0674"] = val
                continue
            if key == "eventKey":
                continue
            rec[imo] = val
        canon.append(rec)
    # 코드리스트 검증
    for k, v in FUEL.items():
        assert v[0] in fuel_codes, f"fuel code {v[0]} not in FAL50 Fuel type list"
    return table, canon


# ── 2. GEARs Type1 변환 ────────────────────────────────────────────────────────────
def template_columns():
    ws = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)["Type 1"]
    cols, grp = [], None
    rows = {r: [c for c in ws.iter_rows(min_row=r, max_row=r, values_only=True)][0] for r in (4, 5, 6, 7)}
    for c in range(len(rows[5])):
        g, n, sb, f = (re.sub(r"\s+", " ", str(rows[r][c])).strip() if rows[r][c] is not None else None for r in (4, 5, 6, 7))
        if g: grp = g
        if n or sb:
            cols.append({"col": openpyxl.utils.get_column_letter(c + 1), "group": grp, "name": n, "sub": sb, "format": f})
    last = None
    for x in cols:
        if x["name"]: last = x["name"]
        x["name_ff"] = x["name"] or last
    return cols


def unlocodes():
    ws = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)["UNLOCODE"]
    s = set()
    for r in ws.iter_rows(min_row=4, values_only=True):  # 시트 A열은 공백, B=Country Code, C=Location code
        if r[1] and r[2] and len(str(r[1])) == 2 and len(str(r[2])) == 3:
            s.add(f"{r[1]}{r[2]}")
    assert len(s) > 10000, f"UNLOCODE sheet parse failed: {len(s)}"
    return s


def hours(a, b):
    f = lambda t: datetime.fromisoformat(t.replace("Z", "+00:00"))
    return round((f(b) - f(a)).total_seconds() / 3600, 2)


def build_legs(canon):
    legs, cur = [], None
    for i, e in enumerate(canon):
        k = e["_event_key"]
        if k == "DEPARTURE_SBY":
            if cur: legs.append(cur)
            cur = {"dep": e, "sea": [], "arr": None, "port": []}
        elif cur and cur["arr"] is None:
            cur["sea"].append(e)
            if k == "ARRIVAL": cur["arr"] = e
        elif cur:
            cur["port"].append(e)
    if cur: legs.append(cur)
    return legs


def rob(e, fuel_col):
    return round(sum(v.get("IMO0674") or 0 for fc, v in e["rob_by_fuel_type"].items() if _gears_fuel(fc) == fuel_col), 3)


def _gears_fuel(fal_code):
    return next(g for p, (c, g) in FUEL.items() if c == fal_code)


def cons(events, fuel_col):
    return round(sum(q for e in events for fc, occ in e["consumption_by_fuel_type"].items() if _gears_fuel(fc) == fuel_col
                     for k, q in occ.items() if k != "IMO0654" and isinstance(q, (int, float))), 3)


def leg_to_row(leg, general, prev_port_events):
    d, a = leg["dep"], leg["arr"]
    fmt_d = lambda t: (t[:10], t[11:16]) if t else (None, None)
    row = {"Voyage No.": str(d.get("IMO0191")), "Port code (DEPARTURE)": d.get("IMO0111"), "Departure Date (UTC)": fmt_d(d["_time"])[0], "Departure time (UTC)": fmt_d(d["_time"])[1],
           "Cargo operation (DEPARTURE)": "Yes" if any(p["_event_key"] == "CARGO_WORK" for p in prev_port_events) else None,
           "Port code (ARRIVAL)": a.get("IMO0108") if a else None, "Arrival Date (UTC)": fmt_d(a["_time"])[0] if a else None, "Arrival time (UTC)": fmt_d(a["_time"])[1] if a else None,
           "Cargo operation (ARRIVAL)": "Yes" if any(p["_event_key"] == "CARGO_WORK" for p in leg["port"]) else None,
           "Time spent at sea (hours)": hours(d["_time"], a["_time"]) if a else None,
           "Distance travelled (nm)": round(sum(e.get("IMO0613") or 0 for e in leg["sea"]), 1) if leg["sea"] else None,
           "Total idle spent time at anchorage (hours)": None, "Cargo carried (Passenger)": None, "Cargo carried (Weight or Volume)": None,
           "Transport Work (Passenger)": None, "Transport Work (Weight or Volume)": None}
    bunk = [e for e in leg["sea"] if e["_event_key"] == "BUNKERING"]
    for f in GEARS_FUELS:
        row[f"{f} ROB at departure (MT)"] = rob(d, f); row[f"{f} ROB at arrival (MT)"] = rob(a, f) if a else None
        row[f"At Berth {f} Consumption (MT)"] = cons(leg["port"], f) if leg["port"] else None
        row[f"{f} Consumption (MT)"] = cons(leg["sea"], f) if leg["sea"] else None
        rq = 0.0
        for b in bunk:
            prev = leg["sea"][leg["sea"].index(b) - 1] if leg["sea"].index(b) > 0 else d
            rq += rob(b, f) - rob(prev, f) + cons([b], f)
        row[f"Bunkering before arrival: Received Quantity ({f})"] = round(rq, 3) if bunk else 0.0
    return row


# ── 3. 검증 ───────────────────────────────────────────────────────────────────────
MANDATORY = ["Voyage No.", "Port code (DEPARTURE)", "Departure Date (UTC)", "Departure time (UTC)", "Cargo operation (DEPARTURE)", "Port code (ARRIVAL)",
             "Arrival Date (UTC)", "Arrival time (UTC)", "Cargo operation (ARRIVAL)", "Time spent at sea (hours)", "Distance travelled (nm)", "Total idle spent time at anchorage (hours)"]
MRV_MANDATORY = ["Cargo carried (Passenger)", "Cargo carried (Weight or Volume)"]


def validate(rows, legs, locodes, canon, code_lists):
    findings = []; checks = 0
    for i, (row, leg) in enumerate(zip(rows, legs)):
        rid = f"row{i+1}"
        for c in MANDATORY:
            checks += 1
            if row.get(c) in (None, ""):
                findings.append({"row": rid, "col": c, "code": "MANDATORY_MISSING", "severity": "FAIL" if leg["arr"] else "INCOMPLETE_LEG",
                                 "detail": "레그 미완료(도착 이벤트 없음)" if not leg["arr"] else "payload 에서 도출 불가"})
        for c in MRV_MANDATORY:
            checks += 1
            if row.get(c) in (None, ""): findings.append({"row": rid, "col": c, "code": "MRV_MANDATORY_MISSING", "severity": "FAIL", "detail": "Noon/Event 페이로드에 화물량 항목 없음"})
        for c in ("Port code (DEPARTURE)", "Port code (ARRIVAL)"):
            checks += 1
            v = row.get(c)
            if v and v not in locodes: findings.append({"row": rid, "col": c, "code": "UNLOCODE_NOT_IN_TEMPLATE_LIST", "severity": "WARNING", "detail": f"{v} — 형식(an5)은 적합, 템플릿 UNLOCODE 시트에 미수록(템플릿 목록 갱신 필요)"})
        for c in ("Departure Date (UTC)", "Arrival Date (UTC)"):
            checks += 1
            v = row.get(c)
            if v and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", v): findings.append({"row": rid, "col": c, "code": "DATE_FORMAT", "severity": "FAIL", "detail": v})
        for c, v in row.items():
            if isinstance(v, (int, float)) and v < 0:
                checks += 1; findings.append({"row": rid, "col": c, "code": "NEGATIVE_VALUE", "severity": "FAIL", "detail": v})
        # 질량수지: ROB_dep + bunkered − cons_sea − ROB_arr ≈ 0
        if leg["arr"]:
            for f in GEARS_FUELS:
                dep, arr, cs, bk = row[f"{f} ROB at departure (MT)"], row[f"{f} ROB at arrival (MT)"], row[f"{f} Consumption (MT)"], row[f"Bunkering before arrival: Received Quantity ({f})"]
                if dep or arr or cs or bk:
                    checks += 1
                    diff = round(dep + bk - cs - arr, 3); tol = max(1.0, 0.005 * max(dep, arr))
                    if abs(diff) > tol: findings.append({"row": rid, "col": f, "code": "MASS_BALANCE", "severity": "WARNING", "detail": f"dep {dep} + bunk {bk} − cons {cs} − arr {arr} = {diff} (tol ±{tol})"})
            # 거리 정합: Σ steaming vs distanceToGo(dep) − distanceToGo(arr)
            checks += 1
            dtg = (leg["dep"].get("IMO0615") or 0) - (leg["arr"].get("IMO0615") or 0)
            if abs(dtg - row["Distance travelled (nm)"]) > 0.05 * max(dtg, 1):
                findings.append({"row": rid, "col": "Distance travelled (nm)", "code": "DISTANCE_CROSSCHECK", "severity": "WARNING", "detail": f"Σsteaming {row['Distance travelled (nm)']} vs ΔdistanceToGo {round(dtg,1)}"})
        # 정박 기간: ROB 감소량 vs 보고된 at-berth 소비량 (미보고 소비 탐지)
        if leg["arr"] and leg["port"]:
            for f in GEARS_FUELS:
                drop = round(rob(leg["arr"], f) - rob(leg["port"][-1], f), 3); rep = row[f"At Berth {f} Consumption (MT)"] or 0
                if drop > 0.5 and abs(drop - rep) > max(0.5, 0.05 * drop):
                    checks += 1
                    findings.append({"row": rid, "col": f"At Berth {f} Consumption (MT)", "code": "AT_BERTH_ROB_DELTA_UNREPORTED", "severity": "WARNING",
                                     "detail": f"정박 중 ROB 감소 {drop} MT 이나 이벤트 보고 소비 {rep} MT (CARGO_WORK 이벤트에 소비량 없음)"})
    # IMO 코드리스트: 이벤트 코드·연료 코드
    for e in canon:
        checks += 1
        if e["IMO0597"] not in code_lists["Event type"]: findings.append({"row": e["_event_key"], "col": "IMO0597", "code": "CODE_VALUE_INVALID", "severity": "FAIL", "detail": e["IMO0597"]})
    return {"checks": checks, "findings": findings, "fail": sum(f["severity"] == "FAIL" for f in findings), "warning": sum(f["severity"] == "WARNING" for f in findings),
            "incomplete": sum(f["severity"] == "INCOMPLETE_LEG" for f in findings)}


# ── 4. KR Systems API (GEARs Post DCS/MRV Voyage Template, Nexawave pApiId=64) 데이터 리스트 ───────────
API_SPEC = ROOT / "k-mds/verification/evidence/C02/manual_S04/kr-systems-api/voyage_template_fields.json"
API_FUEL = {"HFO": "Hfo", "LFO": "Lfo", "MDO": "Mdo", "MGO": "Mgo", "LPG (Propane)": "Lpgp", "LPG (Butane)": "Lpgb", "Methanol": "Methanol", "Ethanol": "Ethanol", "LNG": "Lng", "Other": "Other"}
API_HEAD = {"Voyage No.": "VoyageNo", "Port code (DEPARTURE)": "PortCodeDep", "Departure Date (UTC)": "DepartureDate", "Departure time (UTC)": "DepartureTime",
            "Cargo operation (DEPARTURE)": "DepartureCargoOperYn", "Port code (ARRIVAL)": "PortCodeArr", "Arrival Date (UTC)": "ArrivalDate", "Arrival time (UTC)": "ArrivalTime",
            "Cargo operation (ARRIVAL)": "ArrivalCargoOperYn", "Time spent at sea (hours)": "TimeSpentAtSea", "Distance travelled (nm)": "DistanceTravelled",
            "Total idle spent time at anchorage (hours)": "TotalIdleSpentTime", "Cargo carried (Passenger)": "CargoPassenger", "Cargo carried (Weight or Volume)": "CargoWeight",
            "Transport Work (Passenger)": "TransPassenger", "Transport Work (Weight or Volume)": "TransWeight"}


def type1_to_api(row, general, year="2026", verification_type="002"):
    api = {"DataVerifYear": year, "ImoNo": general.get("imoNo"), "VerificationTypeCode": verification_type}
    for col, prop in API_HEAD.items():
        api[prop] = row.get(col)
    for f, p in API_FUEL.items():
        api[f"{p}Depart"] = row.get(f"{f} ROB at departure (MT)"); api[f"{p}Arrival"] = row.get(f"{f} ROB at arrival (MT)")
        api[f"Berth{p}Consump"] = row.get(f"At Berth {f} Consumption (MT)"); api[f"{p}Consump"] = row.get(f"{f} Consumption (MT)")
        api[f"Bdv{p}"] = row.get(f"Bunkering before arrival: Received Quantity ({f})")
    return api


def validate_api(api_rows):
    spec = json.loads(API_SPEC.read_text("utf-8"))["fields"]
    mand = [s["property"] for s in spec if s["mandatory"]]
    findings = []
    for i, a in enumerate(api_rows):
        rid = f"row{i+1}"
        for p in mand:
            if a.get(p) in (None, ""):
                findings.append({"row": rid, "field": p, "code": "API_MANDATORY_MISSING", "severity": "FAIL"})
        if a.get("ImoNo") and not re.fullmatch(r"\d{7}", str(a["ImoNo"])):
            findings.append({"row": rid, "field": "ImoNo", "code": "IMO_NO_FORMAT", "severity": "FAIL", "detail": f"{a['ImoNo']!r} — 'IMO decimal of the ship' 7자리 아님(테스트 선박 ID)"})
        for p in ("DepartureCargoOperYn", "ArrivalCargoOperYn"):
            if a.get(p) not in (None, "Yes", "No"):
                findings.append({"row": rid, "field": p, "code": "YN_FORMAT", "severity": "FAIL", "detail": a.get(p)})
    return {"mandatory_fields": len(mand), "spec_fields": len(spec), "findings": findings, "fail": len(findings)}


def main(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    payload = json.loads(PAYLOAD.read_text("utf-8"))["data"]
    by, reg, cl = load_codebook(), load_registry(), load_code_lists()
    table, canon = canonicalize(payload, by, reg, cl)
    for k, v in table.items():
        v["gears"] = None
    legs = build_legs(canon)
    rows, prev_port = [], []
    for leg in legs:
        rows.append(leg_to_row(leg, payload["general"], prev_port)); prev_port = leg["port"]
    locodes = unlocodes()
    report = validate(rows, legs, locodes, canon, cl)
    report["event_code_notes"] = EVENT_CODE_NOTE; report["fuel_map"] = {k: {"fal50": v[0], "gears": v[1]} for k, v in FUEL.items()}
    stats = {"payload_fields": len(table), "mapped": sum(1 for v in table.values() if v["imo_id"]), "rules": {}}
    for v in table.values(): stats["rules"][v["rule"]] = stats["rules"].get(v["rule"], 0) + 1
    report["mapping_stats"] = stats
    json.dump(table, open(out / "mapping-table.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    md = ["| Provider 필드 | IMO ID | FAL50 명칭 | 형식 | 규칙 | 비고 |", "|---|---|---|---|---|---|"]
    for k, v in sorted(table.items()): md.append(f"| {k} | {v['imo_id'] or '—'} | {v['imo_name'] or ''} | {v['format'] or ''} | {v['rule']} | {v['note'] or ''} |")
    (out / "mapping-table.md").write_text("\n".join(md), "utf-8")
    json.dump(canon, open(out / "imo-canonical-events.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(rows, open(out / "gears-type1-rows.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    with open(out / "gears-type1-rows.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    api_rows = [type1_to_api(r, payload["general"]) for r in rows]
    api_val = validate_api(api_rows)
    api = {"api": "GEARs Post DCS/MRV Voyage Template (POST v1, Nexawave pApiId=64)", "auth": "Token 헤더(필수) — 미보유, 실제 전송 안 함",
           "contract_status": "FIELD_LIST_CONFIRMED (Nexawave ApiDetail 2026-09-13, 132 fields / 52 mandatory); endpoint URL 미노출; 실제 POST 미실행(승인 필요)",
           "assumptions": {"DataVerifYear": "2026 (이벤트 연도)", "VerificationTypeCode": "002 (IMO DCS) — 001(EU MRV) 병행 여부 사람 결정"},
           "vessel": {"imo_no": payload["general"].get("imoNo"), "ship_name": payload["general"].get("shipName")}, "template_revision": "Type1 Rev.2.1 (2024-12-24)",
           "rows": api_rows, "validation": api_val}
    json.dump(api, open(out / "kr-systems-api-datalist.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    report["api_validation"] = api_val
    json.dump(report, open(out / "validation-report.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(json.dumps({"mapping": stats, "legs": len(legs), "validation": {k: report[k] for k in ("checks", "fail", "warning", "incomplete")}}, ensure_ascii=False, indent=1))
    for f in report["findings"]: print("  -", f)
    for r in rows: print("ROW:", {k: v for k, v in r.items() if v not in (None, 0, 0.0)})
    return 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "k-mds/verification/evidence/C02/manual_S04/gears-transform"))
