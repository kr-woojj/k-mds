"""FAL50 공식 IMO Compendium xlsx → validator Registry 적재 준비 도구.

원본은 read-only로만 연다 (k-mds/data/raw/FAL50 은 참조 전용).

결정론적 파생 규칙 (derivation rules):
  DR-1  'IMO Data Set' sheet의 각 데이터 행을 long-format으로 변환한다.
        Dataset 표시 컬럼(Definition~Format 사이, 값 'X')마다 1 record.
        어느 Dataset에도 표시되지 않은 행은 Dataset 빈 값 1 record.
  DR-2  동일 IMO Data Number의 Business Rules 는 전 occurrence 의
        비어있지 않은 값들의 정렬된 합집합("; " join)으로 통일한다.
        (FAL50 원본에서 IMO0629 가 occurrence 별로 BR 유무가 달라
        validator 의 core-field 일관성 검사에 걸리는 것을 해소.
        내용을 발명하지 않는다 — 원본에 있는 값의 합집합만 사용.)
  DR-3  (number, dataset, path) 중복 occurrence 는 첫 행만 유지한다.
  DR-4  'Code list' sheet 를 code_lists.json 으로 추출한다 (값 검증용).

출력:
  var/registry/derived_fal50.csv        validator import 입력
  var/reference/code_lists.json         code list 값 검증 참조
  var/registry/derivation-manifest.json 원본 해시·규칙·영향 요소 기록

사용:
  uv run python tools/prepare_registry.py [--source PATH] [--import-db PATH]
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = PROJECT_ROOT.parent.parent / "data" / "raw" / "FAL50" / "IMO Compendium.xlsx"  # k-mds 루트
DEFAULT_VALIDATOR = PROJECT_ROOT.parent / "imo-compendium-mapping-validator"
REGISTRY_VERSION = "FAL50"

CSV_HEADER = [
    "IMO Data Number",
    "Data Element",
    "Definition",
    "Format",
    "Code Lists",
    "Business Rules",
    "Dataset",
    "Path",
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract(source: Path) -> tuple[list[dict], list[dict], dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook["IMO Data Set"]
        rows_iter = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows_iter)]
        col = {name.strip().lower().replace(" ", "_"): idx for idx, name in enumerate(header)}
        format_idx = header.index("Format")
        dataset_cols = [
            (idx, header[idx]) for idx in range(col["definition"] + 1, format_idx) if header[idx]
        ]

        raw_records: list[dict] = []
        for row in rows_iter:
            number = row[col["data_number"]]
            if number is None or not str(number).strip():
                continue
            datasets = [
                label for idx, label in dataset_cols
                if row[idx] is not None and str(row[idx]).strip().upper() == "X"
            ]
            raw_records.append(
                {
                    "number": str(number).strip().upper(),
                    "name": _cell(row, col.get("data_element")),
                    "definition": _cell(row, col.get("definition")),
                    "format": _cell(row, col.get("format")),
                    "code_lists": _cell(row, col.get("code_lists")),
                    "business_rules": _cell(row, col.get("business_rules")),
                    "path": _cell(row, col.get("path")),
                    "datasets": datasets,
                }
            )

        # DR-2: element 단위 business rule 합집합
        merged_rules: dict[str, set[str]] = {}
        for record in raw_records:
            rules = merged_rules.setdefault(record["number"], set())
            if record["business_rules"]:
                rules.add(record["business_rules"])
        harmonized = {
            number: "; ".join(sorted(rules)) if rules else ""
            for number, rules in merged_rules.items()
        }
        affected = sorted(
            number
            for number, rules in merged_rules.items()
            if any(r["number"] == number and (r["business_rules"] or "") != harmonized[number]
                   for r in raw_records)
        )

        # DR-1 + DR-3: long-format 변환과 occurrence 중복 제거
        seen: set[tuple[str, str, str]] = set()
        out_rows: list[dict] = []
        for record in raw_records:
            targets = record["datasets"] or [""]
            for dataset in targets:
                key = (record["number"], dataset, record["path"] or "")
                if key in seen:
                    continue
                seen.add(key)
                out_rows.append(
                    {
                        "IMO Data Number": record["number"],
                        "Data Element": record["name"],
                        "Definition": record["definition"],
                        "Format": record["format"],
                        "Code Lists": record["code_lists"],
                        "Business Rules": harmonized[record["number"]],
                        "Dataset": dataset,
                        "Path": record["path"],
                    }
                )

        # DR-4: code list 추출
        code_sheet = workbook["Code list"]
        code_iter = code_sheet.iter_rows(values_only=True)
        next(code_iter)  # header: Code list, Code, Code Name, Code Description
        code_lists: dict[str, list[dict]] = {}
        for row in code_iter:
            if row[0] is None or row[1] is None:
                continue
            code_lists.setdefault(str(row[0]).strip(), []).append(
                {
                    "code": str(row[1]).strip(),
                    "name": str(row[2]).strip() if row[2] is not None else None,
                    "description": str(row[3]).strip() if row[3] is not None else None,
                }
            )

        stats = {
            "source_rows": len(raw_records),
            "derived_records": len(out_rows),
            "distinct_elements": len({r["number"] for r in raw_records}),
            "dataset_columns": [label for _, label in dataset_cols],
            "business_rule_harmonized_elements": affected,
            "code_list_count": len(code_lists),
            "code_value_count": sum(len(v) for v in code_lists.values()),
        }
        code_list_out = [
            {"code_list": name, "values": values} for name, values in sorted(code_lists.items())
        ]
        return out_rows, code_list_out, stats
    finally:
        workbook.close()


def _cell(row, idx):
    if idx is None:
        return ""
    value = row[idx]
    return str(value).strip() if value is not None else ""


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--validator", type=Path, default=DEFAULT_VALIDATOR)
    parser.add_argument("--import-db", type=Path, default=PROJECT_ROOT / "var" / "registry.sqlite3")
    parser.add_argument("--skip-import", action="store_true")
    args = parser.parse_args()

    if not args.source.is_file():
        print(f"error=SOURCE_NOT_FOUND path={args.source}")
        return 2

    source_hash = sha256_file(args.source)
    out_rows, code_lists, stats = extract(args.source)

    registry_dir = PROJECT_ROOT / "var" / "registry"
    reference_dir = PROJECT_ROOT / "var" / "reference"
    registry_dir.mkdir(parents=True, exist_ok=True)
    reference_dir.mkdir(parents=True, exist_ok=True)

    csv_path = registry_dir / "derived_fal50.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_HEADER)
        writer.writeheader()
        writer.writerows(out_rows)

    code_list_path = reference_dir / "code_lists.json"
    code_list_path.write_text(
        json.dumps(
            {"source_sha256": source_hash, "fal_version": REGISTRY_VERSION, "code_lists": code_lists},
            ensure_ascii=False, indent=2, sort_keys=True,
        ),
        encoding="utf-8",
    )

    manifest = {
        "source_path": str(args.source),
        "source_sha256": source_hash,
        "registry_version": REGISTRY_VERSION,
        "derivation_rules": ["DR-1 long-format", "DR-2 business-rule-union",
                             "DR-3 occurrence-dedupe", "DR-4 code-list-extract"],
        "derived_csv_sha256": sha256_file(csv_path),
        "generated_at": datetime.now(UTC).isoformat(),
        "stats": stats,
    }
    manifest_path = registry_dir / "derivation-manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8"
    )

    print(f"sourceHash={source_hash}")
    print(f"derivedRecords={stats['derived_records']}")
    print(f"harmonizedElements={','.join(stats['business_rule_harmonized_elements']) or '-'}")
    print(f"codeLists={stats['code_list_count']} codeValues={stats['code_value_count']}")

    if args.skip_import:
        return 0

    # validator 의 공식 CLI entrypoint 로 적재한다 (재구현 금지 원칙).
    command = [
        "uv", "run", "python", "-m", "app.registry", "import",
        "--file", str(csv_path.resolve()),
        "--version", REGISTRY_VERSION,
        "--db", str(args.import_db.resolve()),
    ]
    result = subprocess.run(
        command, cwd=args.validator, capture_output=True, text=True, timeout=600,
    )
    sys.stdout.write(result.stdout)
    if result.returncode != 0:
        sys.stderr.write(result.stderr)
        print(f"importExitCode={result.returncode}")
        return 1
    print(f"importExitCode={result.returncode}")
    print(f"registryDb={args.import_db}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
